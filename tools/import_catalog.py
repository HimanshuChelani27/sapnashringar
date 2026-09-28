"""Upload Sapna-Catalog.xlsx + the sapna-photos folder to the website in one go.

    backend\\.venv\\Scripts\\python tools\\import_catalog.py                 # live site; checks, then asks
    backend\\.venv\\Scripts\\python tools\\import_catalog.py --check         # only check the file, upload nothing
    backend\\.venv\\Scripts\\python tools\\import_catalog.py --site http://127.0.0.1:8000

Needs: pip install openpyxl httpx. Admin password: typed when asked, or env SAPNA_ADMIN_PASSWORD.
Safe to run again: dresses are matched by code and extras by name and get updated, not duplicated.
Photos are only uploaded to dresses that have none yet.
"""
import argparse
import getpass
import os
import re
import sys
from datetime import date, datetime
from pathlib import Path

import httpx
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
CATS = {"women", "men", "kids", "blouse"}
KINDS = {"jewellery", "pagdi", "umbrella", "dupatta", "other"}
IMG = {".jpg", ".jpeg", ".png", ".webp", ".heic", ".heif"}
CODE = re.compile(r"^[A-Za-z]+[0-9]+$")
PHOTO = re.compile(r"^([A-Za-z]+[0-9]+)(?:[\s_-]*\(?([0-9]+)\)?)?$")  # W001, W001 (2), W001-2, W001_2
SETTINGS = {"upi id": "upi_id", "whatsapp number": "whatsapp_number", "shop address": "shop_address",
            "navratri first night": "navratri_start", "pickup & return rules": "pickup_rules"}


def s(v):
    return "" if v is None else str(v).strip()


def rows(ws, ncols):
    for r, row in enumerate(ws.iter_rows(min_row=3, max_col=ncols, values_only=True), 3):
        if any(s(v) for v in row):
            yield r, row


class Checker:
    def __init__(self):
        self.errors, self.warnings = [], []

    def money(self, v, where, required=False):
        if not s(v):
            if required:
                self.errors.append(f"{where} is empty")
            return 0
        try:
            n = round(float(s(v).replace(",", "").replace("₹", "")))
        except ValueError:
            self.errors.append(f"{where}: '{v}' is not a number")
            return 0
        if n < 0:
            self.errors.append(f"{where}: can't be negative")
        return n

    @staticmethod
    def yes(v, default):
        v = s(v).lower()
        return default if not v else v in ("yes", "y", "true", "1", "haan", "ha")


def scan_photos(folder):
    by_code, stray, qr = {}, [], None
    for f in sorted(folder.iterdir()) if folder.exists() else []:
        if f.suffix.lower() not in IMG:
            continue
        if f.stem.lower().replace(" ", "-") in ("upi-qr", "qr", "upiqr"):
            qr = f
            continue
        m = PHOTO.match(f.stem.strip())
        if m:
            by_code.setdefault(m[1].upper(), []).append((int(m[2] or 0), f.name, f))
        else:
            stray.append(f.name)
    return {k: [f for *_, f in sorted(v)] for k, v in by_code.items()}, stray, qr


def read(xlsx, ck):
    wb = load_workbook(xlsx, data_only=True)
    dresses, seen = [], set()
    for r, (code, name, cat, typ, size, rent, dep, jw, jwp, desc, show) in rows(wb["Dresses"], 11):
        where = f"Dresses row {r}"
        code = s(code).upper()
        if not CODE.match(code):
            ck.errors.append(f"{where}: code '{code}' must be letters then numbers, e.g. W001")
        elif code in seen:
            ck.errors.append(f"{where}: code {code} is used twice")
        seen.add(code)
        if not s(name):
            ck.errors.append(f"{where}: dress name is empty")
        if s(cat).lower() not in CATS:
            ck.errors.append(f"{where}: 'For' must be one of {', '.join(sorted(CATS))} (got '{s(cat)}')")
        has_j = ck.yes(jw, False)
        dresses.append(dict(code=code, name=s(name), gender=s(cat).lower(), type=s(typ), size=s(size),
                            rent=ck.money(rent, f"{where}: rent", required=True),
                            deposit=ck.money(dep, f"{where}: deposit"), jewellery_available=has_j,
                            jewellery_price=ck.money(jwp, f"{where}: jewellery price") if has_j else 0,
                            description=s(desc), active=ck.yes(show, True)))
    extras = []
    for r, (code, name, kind, price, dep, desc) in rows(wb["Extras"], 6):
        where = f"Extras row {r}"
        code = s(code).upper()
        if code and not CODE.match(code):
            ck.errors.append(f"{where}: code '{code}' must be letters then numbers, e.g. J01")
        if not s(name):
            ck.errors.append(f"{where}: name is empty")
        if s(kind).lower() not in KINDS:
            ck.errors.append(f"{where}: type must be one of {', '.join(sorted(KINDS))}")
        extras.append(dict(code=code, name=s(name), kind=s(kind).lower(), price=ck.money(price, f"{where}: price", True),
                           deposit=ck.money(dep, f"{where}: deposit"), description=s(desc)))
    settings = {}
    ws = wb["Shop settings"]
    for label, value in ws.iter_rows(min_row=2, max_col=2, values_only=True):
        key = SETTINGS.get(s(label).lower())
        if not key or not s(value):
            continue
        if key == "navratri_start":
            value = to_date(value, ck)
        elif key == "whatsapp_number":
            value = re.sub(r"\D", "", s(value).removesuffix(".0"))
        settings[key] = s(value)
    return dresses, extras, settings


def to_date(v, ck):
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    for fmt in ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d", "%d.%m.%Y"):
        try:
            return datetime.strptime(s(v), fmt).strftime("%Y-%m-%d")
        except ValueError:
            pass
    ck.errors.append(f"Shop settings: Navratri first night '{v}' should look like 11-10-2026")
    return ""


def form(d):
    return {k: ("true" if v is True else "false" if v is False else str(v)) for k, v in d.items() if k != "code" or v}


def files(field, paths):
    return [(field, (p.name, p.read_bytes(), "application/octet-stream")) for p in paths]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--site", default="https://www.sapnashringar.com")
    ap.add_argument("--xlsx", default=ROOT / "Sapna-Catalog.xlsx", type=Path)
    ap.add_argument("--photos", default=ROOT / "sapna-photos", type=Path)
    ap.add_argument("--check", action="store_true", help="only check, upload nothing")
    ap.add_argument("--yes", action="store_true", help="don't ask before uploading")
    a = ap.parse_args()

    ck = Checker()
    dresses, extras, settings = read(a.xlsx, ck)
    photos, stray, qr = scan_photos(a.photos)
    codes = {d["code"] for d in dresses} | {e["code"] for e in extras if e["code"]}
    for d in dresses:
        if d["code"] not in photos:
            ck.warnings.append(f"{d['code']} {d['name']}: no photos found")
    for code in sorted(set(photos) - codes):
        ck.warnings.append(f"photos for {code} but no {code} in the Excel: {', '.join(p.name for p in photos[code])}")
    for name in stray:
        ck.warnings.append(f"photo name not understood (use CODE (1).jpg): {name}")

    print(f"\nFound {len(dresses)} dresses, {len(extras)} extras, {sum(map(len, photos.values()))} photos"
          f"{', UPI QR' if qr else ''}, {len(settings)} shop settings.")
    for w in ck.warnings:
        print("  ! " + w)
    if ck.errors:
        print("\nPlease fix these in the Excel file, then run again:")
        for e in ck.errors:
            print("  x " + e)
        sys.exit(1)
    print("  Everything looks fine.")
    if a.check:
        return
    if not a.yes and input(f"\nUpload to {a.site}? Type yes: ").strip().lower() != "yes":
        return

    c = httpx.Client(base_url=a.site, timeout=180, follow_redirects=True)
    pw = os.environ.get("SAPNA_ADMIN_PASSWORD") or getpass.getpass("Admin password: ")
    r = c.post("/api/admin/login", json={"password": pw})
    if r.status_code != 200:
        sys.exit("Wrong admin password.")
    H = {"Authorization": "Bearer " + r.json()["token"]}

    def call(method, path, **kw):
        r = c.request(method, path, headers=H, **kw)
        if r.status_code >= 400:
            try:
                detail = r.json().get("detail")
            except ValueError:
                detail = r.text[:200]
            raise SystemExit(f"Stopped: {method} {path} failed ({r.status_code}): {detail}")
        return r.json()

    have = {d["code"].upper(): d for d in call("GET", "/api/admin/dresses") if d["code"]}
    for d in dresses:
        mine = photos.get(d["code"], [])
        old = have.get(d["code"])
        if old:
            call("PUT", f"/api/admin/dresses/{old['id']}", data=form(d))
            added = mine if mine and not old["photos"] else []
            if added:
                call("POST", f"/api/admin/dresses/{old['id']}/photos", files=files("photos", added))
            print(f"  updated  {d['code']} {d['name']}" + (f" (+{len(added)} photos)" if added else ""))
        else:
            call("POST", "/api/admin/dresses", data=form(d), files=files("photos", mine) or None)
            print(f"  added    {d['code']} {d['name']} ({len(mine)} photo{'s' * (len(mine) != 1)})")

    have_x = {x["name"].lower(): x for x in call("GET", "/api/admin/addons")}
    for e in extras:
        pic = photos.get(e["code"], [])[:1]
        old = have_x.get(e["name"].lower())
        body = form({k: v for k, v in e.items() if k != "code"})
        if old:
            call("PUT", f"/api/admin/addons/{old['id']}", data=body, files=files("photo", pic) if pic and not old["photo"] else None)
            print(f"  updated  extra {e['name']}")
        else:
            call("POST", "/api/admin/addons", data=body, files=files("photo", pic) or None)
            print(f"  added    extra {e['name']}")

    if settings or qr:
        call("PUT", "/api/admin/settings", data=settings, files=files("upi_qr", [qr]) if qr else None)
        print(f"  saved    shop settings{' + UPI QR' if qr else ''}")
    print(f"\nDone. Check: {a.site}/garba/home")


if __name__ == "__main__":
    main()
